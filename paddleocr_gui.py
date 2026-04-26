"""Tkinter GUI for running PaddleOCR on images and PDFs."""

from __future__ import annotations

import os
import queue
import threading
import traceback
from dataclasses import dataclass
from typing import Any, Dict, Iterable, List, Optional

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook
from PIL import Image, ImageOps, ImageTk

try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
except Exception:  # Drag and drop is optional.
    DND_FILES = None
    TkinterDnD = None

try:
    from paddleocr import PaddleOCR
except Exception:
    PaddleOCR = None


IMAGE_EXTENSIONS = {
    ".bmp",
    ".jpeg",
    ".jpg",
    ".png",
    ".tif",
    ".tiff",
    ".webp",
}
INPUT_EXTENSIONS = IMAGE_EXTENSIONS | {".pdf"}
LANGUAGE_OPTIONS: Dict[str, str] = {
    "繁體中文（台灣）": "chinese_cht",
    "簡體中文": "ch",
    "英文": "en",
    "日文": "japan",
    "韓文": "korean",
    "拉丁語系": "latin",
    "阿拉伯語系": "arabic",
    "西里爾語系": "cyrillic",
}


@dataclass
class OcrRow:
    file: str
    page: int
    line: int
    text: str
    score: Optional[float]
    box: str


def normalize_paths(paths: Iterable[str]) -> List[str]:
    unique: List[str] = []
    seen = set()
    for raw_path in paths:
        path = os.path.abspath(os.path.expanduser(raw_path.strip()))
        if not os.path.isfile(path):
            continue
        if os.path.splitext(path)[1].lower() not in INPUT_EXTENSIONS:
            continue
        if path not in seen:
            seen.add(path)
            unique.append(path)
    return unique


def value_from_result(result: Any, key: str, default: Any) -> Any:
    if hasattr(result, "get"):
        return result.get(key, default)
    try:
        return result[key]
    except Exception:
        return default


def sequence_from_result(result: Any, key: str) -> Any:
    value = value_from_result(result, key, [])
    return [] if value is None else value


def box_to_string(box: Any) -> str:
    if box is None:
        return ""
    if hasattr(box, "tolist"):
        box = box.tolist()
    return str(box)


class PaddleOcrGui:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("PaddleOCR 圖片文字辨識")
        self.root.geometry("1040x700")
        self.root.minsize(900, 560)

        self.files: List[str] = []
        self.rows: List[OcrRow] = []
        self.preview_image: Optional[ImageTk.PhotoImage] = None
        self.ocr: Optional[Any] = None
        self.ocr_options: Optional[tuple] = None
        self.worker: Optional[threading.Thread] = None
        self.events: "queue.Queue[tuple]" = queue.Queue()

        self.language_var = tk.StringVar(value="繁體中文（台灣）")
        self.orientation_var = tk.BooleanVar(value=False)
        self.status_var = tk.StringVar(value="就緒")

        self._build_ui()
        self._configure_drag_drop()

    def _build_ui(self) -> None:
        outer = ttk.Frame(self.root, padding=12)
        outer.pack(fill=tk.BOTH, expand=True)
        outer.columnconfigure(0, weight=0)
        outer.columnconfigure(1, weight=1)
        outer.rowconfigure(1, weight=1)

        toolbar = ttk.Frame(outer)
        toolbar.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 10))
        toolbar.columnconfigure(8, weight=1)

        self.add_button = ttk.Button(toolbar, text="加入檔案", command=self.add_files)
        self.add_button.grid(row=0, column=0, padx=(0, 6))

        self.clear_button = ttk.Button(toolbar, text="清除", command=self.clear_files)
        self.clear_button.grid(row=0, column=1, padx=(0, 16))

        ttk.Label(toolbar, text="辨識語言").grid(row=0, column=2, padx=(0, 4))
        self.language_box = ttk.Combobox(
            toolbar,
            textvariable=self.language_var,
            width=18,
            state="readonly",
            values=tuple(LANGUAGE_OPTIONS.keys()),
        )
        self.language_box.grid(row=0, column=3, padx=(0, 12))
        self.language_box.bind("<<ComboboxSelected>>", self._reset_ocr)

        self.orientation_check = ttk.Checkbutton(
            toolbar,
            text="偵測文字方向",
            variable=self.orientation_var,
            command=self._reset_ocr,
        )
        self.orientation_check.grid(row=0, column=4, padx=(0, 16))

        self.run_button = ttk.Button(toolbar, text="開始辨識", command=self.run_ocr)
        self.run_button.grid(row=0, column=5, padx=(0, 6))

        self.save_txt_button = ttk.Button(
            toolbar, text="匯出 TXT", command=self.save_txt
        )
        self.save_txt_button.grid(row=0, column=6, padx=(0, 6))

        self.save_xlsx_button = ttk.Button(
            toolbar, text="匯出 Excel", command=self.save_xlsx
        )
        self.save_xlsx_button.grid(row=0, column=7, padx=(0, 6))

        left = ttk.Frame(outer)
        left.grid(row=1, column=0, sticky="nsw", padx=(0, 12))
        left.rowconfigure(1, weight=1)

        ttk.Label(left, text="檔案清單").grid(row=0, column=0, sticky="w")
        list_frame = ttk.Frame(left)
        list_frame.grid(row=1, column=0, sticky="nsew", pady=(4, 8))
        list_frame.rowconfigure(0, weight=1)
        list_frame.columnconfigure(0, weight=1)

        self.file_list = tk.Listbox(list_frame, width=38, height=16)
        self.file_list.grid(row=0, column=0, sticky="nsew")
        self.file_list.bind("<<ListboxSelect>>", self.preview_selected)
        file_scroll = ttk.Scrollbar(
            list_frame, orient=tk.VERTICAL, command=self.file_list.yview
        )
        file_scroll.grid(row=0, column=1, sticky="ns")
        self.file_list.configure(yscrollcommand=file_scroll.set)

        ttk.Label(left, text="預覽").grid(row=2, column=0, sticky="w")
        self.preview_label = ttk.Label(
            left,
            text="尚未選擇圖片",
            anchor=tk.CENTER,
            relief=tk.SOLID,
            width=38,
        )
        self.preview_label.grid(row=3, column=0, sticky="ew", pady=(4, 0), ipady=80)

        right = ttk.Frame(outer)
        right.grid(row=1, column=1, sticky="nsew")
        right.rowconfigure(1, weight=1)
        right.columnconfigure(0, weight=1)

        ttk.Label(right, text="辨識結果").grid(row=0, column=0, sticky="w")
        text_frame = ttk.Frame(right)
        text_frame.grid(row=1, column=0, sticky="nsew", pady=(4, 8))
        text_frame.rowconfigure(0, weight=1)
        text_frame.columnconfigure(0, weight=1)

        self.output = tk.Text(text_frame, wrap=tk.WORD, undo=False)
        self.output.grid(row=0, column=0, sticky="nsew")
        output_scroll = ttk.Scrollbar(
            text_frame, orient=tk.VERTICAL, command=self.output.yview
        )
        output_scroll.grid(row=0, column=1, sticky="ns")
        self.output.configure(yscrollcommand=output_scroll.set)

        bottom = ttk.Frame(outer)
        bottom.grid(row=2, column=0, columnspan=2, sticky="ew")
        bottom.columnconfigure(0, weight=1)

        self.progress = ttk.Progressbar(bottom, mode="determinate")
        self.progress.grid(row=0, column=0, sticky="ew", padx=(0, 10))

        ttk.Label(bottom, textvariable=self.status_var).grid(row=0, column=1)

    def _configure_drag_drop(self) -> None:
        if TkinterDnD is None or DND_FILES is None:
            return
        for widget in (self.root, self.file_list):
            try:
                widget.drop_target_register(DND_FILES)
                widget.dnd_bind("<<Drop>>", self.on_drop)
            except Exception:
                continue

    def _reset_ocr(self, *_args: Any) -> None:
        self.ocr = None
        self.ocr_options = None

    def add_files(self) -> None:
        paths = filedialog.askopenfilenames(
            title="選擇圖片或 PDF",
            filetypes=(
                ("圖片與 PDF", "*.png *.jpg *.jpeg *.bmp *.tif *.tiff *.webp *.pdf"),
                ("圖片", "*.png *.jpg *.jpeg *.bmp *.tif *.tiff *.webp"),
                ("PDF", "*.pdf"),
                ("所有檔案", "*.*"),
            ),
        )
        self.add_paths(paths)

    def on_drop(self, event: Any) -> None:
        paths = self.root.tk.splitlist(event.data)
        self.add_paths(paths)

    def add_paths(self, paths: Iterable[str]) -> None:
        new_paths = normalize_paths(paths)
        existing = set(self.files)
        added = [path for path in new_paths if path not in existing]
        if not added:
            return
        self.files.extend(added)
        self.refresh_file_list()
        if len(self.files) == len(added):
            self.file_list.selection_set(0)
            self.preview_selected()
        self.status_var.set(f"已選擇 {len(self.files)} 個檔案")

    def clear_files(self) -> None:
        if self.is_busy():
            return
        self.files.clear()
        self.rows.clear()
        self.file_list.delete(0, tk.END)
        self.output.delete("1.0", tk.END)
        self.preview_image = None
        self.preview_label.configure(image="", text="尚未選擇圖片")
        self.progress.configure(value=0, maximum=1)
        self.status_var.set("就緒")

    def refresh_file_list(self) -> None:
        self.file_list.delete(0, tk.END)
        for path in self.files:
            self.file_list.insert(tk.END, os.path.basename(path))

    def preview_selected(self, *_args: Any) -> None:
        selection = self.file_list.curselection()
        if not selection:
            return
        path = self.files[selection[0]]
        if os.path.splitext(path)[1].lower() not in IMAGE_EXTENSIONS:
            self.preview_image = None
            self.preview_label.configure(image="", text="PDF 不顯示預覽")
            return
        try:
            image = Image.open(path)
            image = ImageOps.exif_transpose(image)
            image.thumbnail((330, 260))
            self.preview_image = ImageTk.PhotoImage(image)
            self.preview_label.configure(image=self.preview_image, text="")
        except Exception as exc:
            self.preview_image = None
            self.preview_label.configure(image="", text=f"預覽失敗：{exc}")

    def is_busy(self) -> bool:
        return self.worker is not None and self.worker.is_alive()

    def set_busy(self, busy: bool) -> None:
        state = tk.DISABLED if busy else tk.NORMAL
        for widget in (
            self.add_button,
            self.clear_button,
            self.run_button,
            self.save_txt_button,
            self.save_xlsx_button,
            self.language_box,
            self.orientation_check,
        ):
            widget.configure(state=state)

    def ensure_ocr(self, options: tuple) -> Any:
        if PaddleOCR is None:
            raise RuntimeError(
                "尚未安裝 paddleocr。請執行：pip install paddleocr paddlepaddle"
            )
        if self.ocr is not None and self.ocr_options == options:
            return self.ocr
        lang, use_orientation = options
        self.events.put(("status", "正在載入 PaddleOCR 模型..."))
        self.ocr = PaddleOCR(
            lang=lang,
            use_doc_orientation_classify=False,
            use_doc_unwarping=False,
            use_textline_orientation=use_orientation,
        )
        self.ocr_options = options
        return self.ocr

    def run_ocr(self) -> None:
        if self.is_busy():
            return
        if not self.files:
            messagebox.showinfo("PaddleOCR", "請先加入至少一個圖片或 PDF。")
            return

        self.rows.clear()
        self.output.delete("1.0", tk.END)
        self.progress.configure(value=0, maximum=len(self.files))
        self.set_busy(True)
        self.status_var.set("開始辨識...")

        files = list(self.files)
        language_code = LANGUAGE_OPTIONS.get(self.language_var.get(), "chinese_cht")
        options = (language_code, bool(self.orientation_var.get()))
        self.worker = threading.Thread(
            target=self._ocr_worker, args=(files, options), daemon=True
        )
        self.worker.start()
        self.root.after(100, self.poll_events)

    def _ocr_worker(self, files: List[str], options: tuple) -> None:
        try:
            ocr = self.ensure_ocr(options)
            rows: List[OcrRow] = []
            for file_index, path in enumerate(files, start=1):
                self.events.put(("status", f"正在辨識 {os.path.basename(path)}"))
                results = ocr.predict(path)
                for result in results:
                    page_index = value_from_result(result, "page_index", None)
                    page = int(page_index) + 1 if page_index is not None else 1
                    texts = sequence_from_result(result, "rec_texts")
                    scores = sequence_from_result(result, "rec_scores")
                    boxes = sequence_from_result(result, "rec_boxes")
                    for line_index, text in enumerate(texts, start=1):
                        score = scores[line_index - 1] if line_index - 1 < len(scores) else None
                        box = boxes[line_index - 1] if line_index - 1 < len(boxes) else None
                        rows.append(
                            OcrRow(
                                file=path,
                                page=page,
                                line=line_index,
                                text=str(text),
                                score=float(score) if score is not None else None,
                                box=box_to_string(box),
                            )
                        )
                self.events.put(("progress", file_index))
            self.events.put(("done", rows))
        except Exception as exc:
            self.events.put(("error", str(exc), traceback.format_exc()))

    def poll_events(self) -> None:
        while True:
            try:
                event = self.events.get_nowait()
            except queue.Empty:
                break
            kind = event[0]
            if kind == "status":
                self.status_var.set(event[1])
            elif kind == "progress":
                self.progress.configure(value=event[1])
            elif kind == "done":
                self.rows = event[1]
                self.display_rows()
                self.set_busy(False)
                self.status_var.set(f"完成，辨識到 {len(self.rows)} 行文字。")
                self.worker = None
            elif kind == "error":
                self.set_busy(False)
                self.worker = None
                self.status_var.set("OCR 失敗")
                self.show_error(event[1], event[2])

        if self.is_busy():
            self.root.after(100, self.poll_events)

    def display_rows(self) -> None:
        self.output.delete("1.0", tk.END)
        if not self.rows:
            self.output.insert(tk.END, "未辨識到文字。")
            return

        current_file = None
        current_page = None
        for row in self.rows:
            if row.file != current_file:
                current_file = row.file
                current_page = None
                self.output.insert(tk.END, f"\n[{os.path.basename(row.file)}]\n")
            if row.page != current_page:
                current_page = row.page
                self.output.insert(tk.END, f"第 {row.page} 頁\n")
            self.output.insert(tk.END, row.text + "\n")
        self.output.see("1.0")

    def show_error(self, message: str, details: str) -> None:
        self.output.delete("1.0", tk.END)
        self.output.insert(tk.END, details)
        messagebox.showerror("OCR 失敗", message)

    def save_txt(self) -> None:
        if not self.rows:
            messagebox.showinfo("PaddleOCR", "請先執行辨識再匯出。")
            return
        path = filedialog.asksaveasfilename(
            title="匯出文字檔",
            defaultextension=".txt",
            filetypes=(("文字檔", "*.txt"), ("所有檔案", "*.*")),
        )
        if not path:
            return
        with open(path, "w", encoding="utf-8") as handle:
            current_file = None
            current_page = None
            for row in self.rows:
                if row.file != current_file:
                    current_file = row.file
                    current_page = None
                    handle.write(f"\n[{os.path.basename(row.file)}]\n")
                if row.page != current_page:
                    current_page = row.page
                    handle.write(f"第 {row.page} 頁\n")
                handle.write(row.text + "\n")
        self.status_var.set(f"已匯出：{path}")

    def save_xlsx(self) -> None:
        if not self.rows:
            messagebox.showinfo("PaddleOCR", "請先執行辨識再匯出。")
            return
        path = filedialog.asksaveasfilename(
            title="匯出 Excel 活頁簿",
            defaultextension=".xlsx",
            filetypes=(("Excel 活頁簿", "*.xlsx"), ("所有檔案", "*.*")),
        )
        if not path:
            return

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "辨識結果"
        sheet.append(("檔案", "頁數", "行號", "文字", "信心分數", "位置"))
        for row in self.rows:
            sheet.append(
                (
                    row.file,
                    row.page,
                    row.line,
                    row.text,
                    row.score,
                    row.box,
                )
            )
        workbook.save(path)
        self.status_var.set(f"已匯出：{path}")


def create_root() -> tk.Tk:
    if TkinterDnD is not None:
        return TkinterDnD.Tk()
    return tk.Tk()


def main() -> None:
    root = create_root()
    app = PaddleOcrGui(root)
    root.mainloop()


if __name__ == "__main__":
    main()
